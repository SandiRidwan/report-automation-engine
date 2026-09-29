"""
test_data_quality.py — uji kualitas data & laporan (CI-friendly).
"""
from __future__ import annotations

import glob
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

import duckdb

from config import COUNTRIES, DB_FILE, INDICATORS, OUTPUT, STAGING

FAILS: list[str] = []
PASSES: list[str] = []


def check(name, cond, detail=""):
    if cond:
        PASSES.append(name)
        print(f"  PASS  {name}")
    else:
        FAILS.append(f"{name} — {detail}")
        print(f"  FAIL  {name}  {detail}")


def main() -> int:
    print("[dq] uji kualitas data & laporan Excel\n")

    # --- panel data ---
    import pandas as pd
    panel = STAGING / "wb_panel.parquet"
    check("panel data ada", panel.exists())
    if panel.exists():
        df = pd.read_parquet(panel)
        check("panel tidak kosong", len(df) > 0, f"{len(df)}")
        check("semua nilai terisi", df["value"].notna().all())
        check("cakupan negara", df.country_iso.nunique() == len(COUNTRIES),
              f"{df.country_iso.nunique()}")
        check("cakupan indikator",
              df.indicator_code.nunique() == len(INDICATORS),
              f"{df.indicator_code.nunique()}")
        check("tahun dalam rentang wajar",
              df["year"].between(2000, 2100).all())

    # --- marts ---
    con = duckdb.connect(str(DB_FILE), read_only=True)
    for t in ["mart_latest", "mart_growth", "mart_category_avg"]:
        n = con.execute(f"SELECT count(*) FROM {t}").fetchone()[0]
        check(f"{t} terisi", n > 0, f"{n}")
    # growth punya null (jika v0=0) tapi tidak semuanya null
    g = con.execute("SELECT count(*) FROM mart_growth WHERE growth_pct IS NOT NULL").fetchone()[0]
    check("growth_pct terhitung", g > 0, f"{g}")
    con.close()

    # --- laporan Excel ---
    files = sorted(glob.glob(str(OUTPUT / "*.xlsx")))
    check("laporan Excel dihasilkan", len(files) > 0, f"{len(files)} file")
    if files:
        from openpyxl import load_workbook
        wb = load_workbook(files[-1])
        check("5 sheet", len(wb.sheetnames) == 5, str(wb.sheetnames))
        nch = sum(len(wb[s]._charts) for s in wb.sheetnames)
        check("chart ada", nch > 0, f"{nch}")

    print(f"\n[dq] {len(PASSES)} lulus, {len(FAILS)} gagal")
    if FAILS:
        print("\nGAGAL:")
        for f in FAILS:
            print("  -", f)
        return 1
    print("[dq] SEMUA UJI LULUS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
